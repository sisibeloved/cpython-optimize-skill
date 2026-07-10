#!/usr/bin/env python3
"""
分析指定的 pyperformance JSON 结果文件，
生成 Excel 对比表格和分页性能趋势图（每20个用例一张图）。

用法：
  python get_stat.py a.json d.json c.json b.json
  python get_stat.py -c a.json d.json c.json b.json
  python get_stat.py -b 2to3 -b chaos -b nbody a.json d.json c.json b.json
  python get_stat.py -c -b 2to3 -b chaos a.json b.json
  python get_stat.py

依赖：-c/--console-only 不需要额外依赖；生成 Excel/PNG 需要 pip install openpyxl matplotlib
"""

import json
import glob
import os
import argparse
from math import exp, log
from pathlib import Path


def import_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise SystemExit(
            "Missing dependency: openpyxl. Install with: python -m pip install openpyxl"
        ) from exc
    return Workbook, Font, Alignment, PatternFill, Border, Side, get_column_letter


def import_pyplot():
    try:
        import matplotlib
        import matplotlib.pyplot as plt
    except Exception as exc:
        raise SystemExit(
            "Missing dependency: matplotlib. Install with: python -m pip install matplotlib"
        ) from exc

    matplotlib.rcParams['font.sans-serif'] = [
        'SimHei', 'WenQuanYi Micro Hei', 'Noto Sans CJK SC',
        'Microsoft YaHei', 'DejaVu Sans',
    ]
    matplotlib.rcParams['axes.unicode_minus'] = False
    return plt


# ═══════════════════════════════════════════════════════════
#  工具函数
# ═══════════════════════════════════════════════════════════

def load_benchmark_data(json_file):
    """从 pyperformance JSON 中提取各用例平均耗时(秒)。"""
    with open(json_file, 'r', encoding='utf-8') as f:
        data = json.load(f)

    results = {}
    if 'benchmarks' in data:
        for bench in data['benchmarks']:
            name = None
            if 'metadata' in bench and 'name' in bench['metadata']:
                name = bench['metadata']['name']
            if name is None:
                for run in bench.get('runs', []):
                    if 'metadata' in run and 'name' in run['metadata']:
                        name = run['metadata']['name']
                        break
            if name is None:
                continue
            values = []
            for run in bench.get('runs', []):
                if 'values' in run:
                    values.extend(run['values'])
            if values:
                results[name] = sum(values) / len(values)
    elif 'runs' in data:
        name = data.get('metadata', {}).get('name', 'unknown')
        values = []
        for run in data['runs']:
            if 'values' in run:
                values.extend(run['values'])
        if values:
            results[name] = sum(values) / len(values)
    return results


def choose_unit(value_seconds):
    """根据数值量级选取合适的显示单位。"""
    if value_seconds >= 1.0:
        return 'sec', 1.0
    elif value_seconds >= 1e-3:
        return 'ms', 1e-3
    elif value_seconds >= 1e-6:
        return 'us', 1e-6
    else:
        return 'ns', 1e-9


def geometric_mean(values):
    """计算一组正数的几何平均值。"""
    pos = [v for v in values if v > 0]
    if not pos:
        return 1.0
    return exp(sum(log(v) for v in pos) / len(pos))


def make_unique_file_labels(file_paths):
    """为输入文件生成最短且唯一的路径后缀标签。"""
    path_parts = [Path(path).resolve().parts for path in file_paths]
    labels = []

    for parts in path_parts:
        label = os.path.join(*parts)
        for depth in range(1, len(parts) + 1):
            suffix = parts[-depth:]
            matches = sum(
                1
                for other in path_parts
                if len(other) >= depth and other[-depth:] == suffix
            )
            if matches == 1:
                label = os.path.join(*suffix)
                break
        labels.append(label)

    totals = {label: labels.count(label) for label in labels}
    seen = {}
    unique_labels = []
    for label in labels:
        if totals[label] == 1:
            unique_labels.append(label)
            continue
        seen[label] = seen.get(label, 0) + 1
        unique_labels.append(f"{label}#{seen[label]}")
    return unique_labels


def remove_stale_plot_files(base_name='benchmark_trends'):
    """删除脚本拥有的旧分页 PNG，避免重跑后混入过期页面。"""
    base_path = Path(base_name)
    directory = base_path.parent
    prefix = f"{base_path.name}_part"
    for path in directory.glob(f"{prefix}*.png"):
        page_number = path.stem[len(prefix):]
        if page_number.isdigit():
            path.unlink()


# ═══════════════════════════════════════════════════════════
#  Excel 输出
# ═══════════════════════════════════════════════════════════

def save_excel(common_benchmarks, valid_files, file_labels, all_data, units, divisors,
               perf_geo_means, excel_support, xlsx_path='benchmark_comparison.xlsx'):
    """生成带格式的 Excel 对比表格。"""
    Workbook, Font, Alignment, PatternFill, Border, Side, get_column_letter = excel_support

    wb = Workbook()
    ws = wb.active
    ws.title = '性能对比'

    header_font = Font(bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    summary_font = Font(bold=True, size=11, color='FFFFFF')
    summary_fill = PatternFill(start_color='E26B0A', end_color='E26B0A', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    center = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    headers = ['Benchmark', 'Unit'] + file_labels
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border

    for row_idx, bench in enumerate(common_benchmarks, 2):
        c = ws.cell(row=row_idx, column=1, value=bench)
        c.alignment = left_align
        c.border = thin_border

        c = ws.cell(row=row_idx, column=2, value=units[bench])
        c.alignment = center
        c.border = thin_border

        for col_off, jf in enumerate(valid_files):
            v = all_data[jf][bench] / divisors[bench]
            c = ws.cell(row=row_idx, column=3 + col_off, value=round(v, 4))
            c.number_format = '0.0000'
            c.alignment = center
            c.border = thin_border

    summary_row = len(common_benchmarks) + 2
    c = ws.cell(row=summary_row, column=1, value='性能对比')
    c.font = summary_font
    c.fill = summary_fill
    c.alignment = left_align
    c.border = thin_border

    c = ws.cell(row=summary_row, column=2, value='NA')
    c.font = summary_font
    c.fill = summary_fill
    c.alignment = center
    c.border = thin_border

    for col_off, gm in enumerate(perf_geo_means):
        c = ws.cell(row=summary_row, column=3 + col_off, value=round(gm, 4))
        c.font = summary_font
        c.fill = summary_fill
        c.number_format = '0.0000'
        c.alignment = center
        c.border = thin_border

    ws.column_dimensions['A'].width = max(len(b) for b in common_benchmarks) + 4
    ws.column_dimensions['B'].width = 8
    for col_off, lb in enumerate(file_labels):
        col_letter = get_column_letter(3 + col_off)
        ws.column_dimensions[col_letter].width = max(len(lb), 12) + 4

    ws.freeze_panes = 'A2'
    wb.save(xlsx_path)
    print(f"📄  Excel 表格已保存到 {xlsx_path}")


# ═══════════════════════════════════════════════════════════
#  绘图（分页）
# ═══════════════════════════════════════════════════════════

def plot_trends_paginated(common_benchmarks, valid_files, file_labels, perf_ratios,
                          perf_geo_means, plt, benchmarks_per_page=20,
                          base_name='benchmark_trends'):
    n_bench = len(common_benchmarks)
    n_pages = (n_bench + benchmarks_per_page - 1) // benchmarks_per_page
    x = list(range(len(valid_files)))

    print(f"\n📊  共 {n_bench} 个用例，将生成 {n_pages} 张图（每页最多 {benchmarks_per_page} 个用例）")

    for page_idx in range(n_pages):
        start = page_idx * benchmarks_per_page
        end = min(start + benchmarks_per_page, n_bench)
        page_benchmarks = common_benchmarks[start:end]
        n_plots = len(page_benchmarks) + 1

        subplot_h = 2.0
        fig_h = n_plots * subplot_h
        fig_w = max(10, len(valid_files) * 1.5)

        fig, axes = plt.subplots(
            n_plots, 1,
            figsize=(fig_w, fig_h),
            sharex=True,
        )
        if n_plots == 1:
            axes = [axes]

        for idx, bench in enumerate(page_benchmarks):
            ax = axes[idx]
            y = [perf_ratios[jf][bench] for jf in valid_files]
            ax.plot(x, y, 'o-', color='steelblue', linewidth=1.8, markersize=6)
            ax.fill_between(x, 1.0, y, alpha=0.10, color='steelblue')
            ax.axhline(y=1.0, color='red', linestyle='--', alpha=0.5, linewidth=0.8)

            for i, v in enumerate(y):
                ax.annotate(f'{v:.3f}', (i, v),
                            textcoords='offset points', xytext=(0, 8),
                            ha='center', fontsize=7, color='steelblue')

            ax.set_ylabel('Ratio', fontsize=8)
            ax.set_title(bench, fontsize=10, fontweight='bold', loc='left', pad=4)
            ax.grid(True, alpha=0.2)
            ax.tick_params(labelsize=8)

        ax = axes[-1]
        color = 'forestgreen' if perf_geo_means[-1] >= 1.0 else 'orangered'
        ax.plot(x, perf_geo_means, 'o-', color=color, linewidth=2.8, markersize=9)
        ax.fill_between(x, 1.0, perf_geo_means, alpha=0.18, color=color)
        ax.axhline(y=1.0, color='red', linestyle='--', alpha=0.5, linewidth=0.8)

        for i, gm in enumerate(perf_geo_means):
            ax.annotate(f'{gm:.4f}', (i, gm),
                        textcoords='offset points', xytext=(0, 12),
                        ha='center', fontsize=10, fontweight='bold', color=color)

        ax.set_ylabel('Geo Mean', fontsize=9)
        ax.set_title('★ Overall Performance (Geometric Mean)', fontsize=11,
                     fontweight='bold', loc='left', color=color, pad=4)
        ax.grid(True, alpha=0.2)

        ax.set_xticks(x)
        ax.set_xticklabels(file_labels, rotation=45, ha='right', fontsize=9)
        ax.set_xlabel('JSON Files', fontsize=10)

        png_path = f"{base_name}_part{page_idx + 1}.png"
        plt.tight_layout(h_pad=0.8)
        plt.savefig(png_path, dpi=150, bbox_inches='tight')
        print(f"   ✅ {png_path}  ({len(page_benchmarks)} 个用例)")
        plt.close()

    print()


# ═══════════════════════════════════════════════════════════
#  参数解析
# ═══════════════════════════════════════════════════════════

def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description='pyperformance 多版本性能对比分析工具',
        epilog='示例:\n'
               '  python get_stat.py a.json d.json c.json b.json\n'
               '  python get_stat.py -c a.json d.json c.json b.json\n'
               '  python get_stat.py -b 2to3 -b chaos -b nbody a.json b.json\n'
               '  python get_stat.py -c -b 2to3 -b chaos a.json b.json\n'
               '  python get_stat.py\n',
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        '-c', '--console-only',
        action='store_true',
        help='仅打屏输出对比数据，不生成 Excel 和趋势图',
    )
    parser.add_argument(
        '-b', '--benchmarks',
        action='append',
        metavar='NAME',
        help='只查看指定的用例；指定多个用例时重复使用 -b，不指定则查看全部公共用例',
    )
    parser.add_argument(
        'json_files',
        nargs='*',
        help='要对比的 JSON 文件（按指定顺序）；不传则自动扫描当前目录',
    )
    return parser.parse_args(argv)


# ═══════════════════════════════════════════════════════════
#  主流程
# ═══════════════════════════════════════════════════════════

def main(argv=None):
    args = parse_args(argv)

    # ── 1. 获取文件列表 ──
    if args.json_files:
        json_files = args.json_files
        for jf in json_files:
            if not os.path.isfile(jf):
                print(f"❌  文件不存在: {jf}")
                return 1
    else:
        json_files = sorted(glob.glob('*.json'))

    if not json_files:
        print("❌  未找到 JSON 文件。")
        print(f"用法: python {os.path.basename(__file__)} [-c] [-b 用例]... a.json b.json")
        return 1

    if len(json_files) < 2:
        print("❌  至少需要两个 JSON 文件：第一个是 baseline，后续文件是 candidate。")
        return 1

    print(f"📂  共 {len(json_files)} 个 JSON 文件（按指定顺序）：")
    baseline_file = json_files[0]
    all_data = {}

    try:
        baseline_data = load_benchmark_data(baseline_file)
    except Exception as e:
        print(f"   ❌  baseline {baseline_file} 加载失败: {e}")
        return 1
    if not baseline_data:
        print(f"   ❌  baseline {baseline_file} 没有可用 benchmark 数据。")
        return 1

    all_data[baseline_file] = baseline_data
    valid_files = [baseline_file]
    print(f"   ✅ {baseline_file}  ({len(baseline_data)} 个用例，baseline)")

    for jf in json_files[1:]:
        try:
            candidate_data = load_benchmark_data(jf)
        except Exception as e:
            print(f"   ⚠️  {jf} 加载失败，已跳过: {e}")
            continue
        if not candidate_data:
            print(f"   ⚠️  {jf} 没有可用 benchmark 数据，已跳过")
            continue
        all_data[jf] = candidate_data
        valid_files.append(jf)
        print(f"   ✅ {jf}  ({len(candidate_data)} 个用例)")

    if len(valid_files) < 2:
        print("❌  除有效 baseline 外，至少还需要一个有效 candidate JSON。")
        return 1

    # ── 2. 取公共用例 ──
    common = set(all_data[baseline_file].keys())
    for jf in valid_files[1:]:
        common &= set(all_data[jf].keys())

    # ── 3. 如果指定了 -b，只保留指定的用例 ──
    if args.benchmarks:
        selected = []
        not_found = []
        for b in args.benchmarks:
            if b in common:
                selected.append(b)
            else:
                not_found.append(b)
        if not_found:
            print(f"\n⚠️  以下用例在公共用例中未找到，已跳过: {', '.join(not_found)}")
            print(f"   可用的公共用例共 {len(common)} 个：")
            for name in sorted(common):
                print(f"     - {name}")
        if not selected:
            print("❌  指定的用例均不存在。")
            return 1
        common_benchmarks = selected
        print(f"\n🎯  已筛选 {len(common_benchmarks)} 个指定用例: {', '.join(common_benchmarks)}\n")
    else:
        common_benchmarks = sorted(common)

    if not common_benchmarks:
        print("❌  多个文件之间没有公共用例。")
        return 1

    if not args.benchmarks:
        print(f"\n🔗  公共用例数量: {len(common_benchmarks)}\n")

    # ── 4. 确定显示单位 ──
    units = {}
    divisors = {}
    for bench in common_benchmarks:
        u, d = choose_unit(all_data[baseline_file][bench])
        units[bench] = u
        divisors[bench] = d

    # ── 5. 计算性能对比 ──
    perf_ratios = {}
    perf_geo_means = []
    for jf in valid_files:
        ratios = {}
        for bench in common_benchmarks:
            ratios[bench] = all_data[baseline_file][bench] / all_data[jf][bench]
        perf_ratios[jf] = ratios
        perf_geo_means.append(geometric_mean(list(ratios.values())))

    # ── 6. 终端表格（始终输出） ──
    file_labels = make_unique_file_labels(valid_files)
    col_w = max(max(len(label) for label in file_labels), 12)
    name_w = max(max(len(bench) for bench in common_benchmarks), 8)
    unit_w = 5

    header = f"{'Benchmark':<{name_w}}  {'Unit':<{unit_w}}"
    for label in file_labels:
        header += f"  {label:>{col_w}}"
    sep = '=' * len(header)
    thin = '-' * len(header)

    print(sep)
    print(header)
    print(sep)
    for bench in common_benchmarks:
        row = f"{bench:<{name_w}}  {units[bench]:<{unit_w}}"
        for jf in valid_files:
            value = all_data[jf][bench] / divisors[bench]
            row += f"  {value:>{col_w}.4f}"
        print(row)
    print(thin)
    row = f"{'性能对比':<{name_w}}  {'NA':<{unit_w}}"
    for geometric_mean_value in perf_geo_means:
        row += f"  {geometric_mean_value:>{col_w}.4f}"
    print(row)
    print(sep)

    # ── 7. 如果指定了 -c，到此结束 ──
    if args.console_only:
        print("\n💡  已指定 -c 参数，仅打屏输出，跳过 Excel 和趋势图生成。")
        return 0

    # ── 8. 在写入任何报告前预检依赖并清理旧分页 ──
    excel_support = import_openpyxl()
    plt = import_pyplot()
    remove_stale_plot_files()

    # ── 9. 保存 Excel ──
    save_excel(common_benchmarks, valid_files, file_labels, all_data, units, divisors,
               perf_geo_means, excel_support)

    # ── 10. 分页绘图 ──
    plot_trends_paginated(common_benchmarks, valid_files, file_labels, perf_ratios,
                          perf_geo_means, plt, benchmarks_per_page=20)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
